from docx import Document
import xlrd

document = Document()
table = document.add_table(rows=50,cols=5,style='Light Grid')
wb = xlrd.open_workbook('111.xlsx')
sheetname = wb.sheet_names()[0]
sheet_c = wb.sheet_by_index(0)

k=7
j=0
old=sheet_c.row_values(1)[0]
print("old:"+old)
for i in range(1, sheet_c.nrows):
    row_data = sheet_c.row_values(i)
    # add_row = table.add_row().cells
    # table.cell(i,1).text=row_data[0]
    newv=row_data[0]
    print('newv:'+newv)

#当不相等的时候，我们需要切开案例
    if old!=newv:
        j=j+11
        print('j:'+str(j))
        add_row = table.add_row().cells

#需要在这里构造出报表

        table.cell(i+j+k-7,0).text=u'用例名称'
        table.cell(i+j+k-7,1).text=row_data[1]

        table.cell(i+j+k-7,2).text=u'用例标识'
        table.cell(i+j+k-7,3).merge(table.cell(i+j+k-7,4)) #合并单元格
        table.cell(i+j+k-7,3).text=row_data[0]

        table.cell(i+j+k-6,0).text=u'用例追踪'
        table.cell(i+j+k-6,1).merge(table.cell(i+j+k-6,4))
        table.cell(i+j+k-6,1).text=row_data[2]

        table.cell(i+j+k-5,0).text=u'用例说明'
        table.cell(i+j+k-5,1).merge(table.cell(i+j+k-5,4))
        table.cell(i+j+k-5,1).text=row_data[3]

        table.cell(i+j+k-4,0).text=u'用例初始化'
        table.cell(i+j+k-4,1).merge(table.cell(i+j+k-4,4))
        table.cell(i+j+k-4,1).text=row_data[4]

        table.cell(i+j+k-3,0).text=u'测试过程'
        table.cell(i+j+k-3,0).merge(table.cell(i+j+k-3,4))

        table.cell(i+j+k-2,0).text=u'前提及约束'
        table.cell(i+j+k-2,1).merge(table.cell(i+j+k-2,4))
        table.cell(i+j+k-2,1).text=row_data[4]

        table.cell(i+j+k-1,0).text=u'序号'
        table.cell(i+j+k,0).text=row_data[6]

        table.cell(i+j+k-1,1).text=u'输入及操作说明'
        table.cell(i+j+k,1).text=row_data[7]

        table.cell(i+j+k-1,2).text=u'期望'
        table.cell(i+j+k,2).text=row_data[8]

        table.cell(i+j+k-1,3).text=u'评估标准'
        table.cell(i+j+k,3).text=row_data[9]

        table.cell(i+j+k-1,4).text=u'备注'
        table.cell(i+j+k,4).text=row_data[10]

        table.cell(i+j+k-11,0).text=u'测试终止条件'
        table.cell(i+j+k-11,1).merge(table.cell(i+j+k-11,4))
        table.cell(i+j+k-11,1).text=u'测试结果符合预期后，测试终止'

        table.cell(i+j+k-10,0).text=u'测试结果'
        table.cell(i+j+k-10,1).merge(table.cell(i+j+k-10,4))
        table.cell(i+j+k-10,1).text=u'□通过  □未通过但可作优化或修改  □未通过且无法修改'

        table.cell(i+j+k-9,0).text=u'设计人员'
        table.cell(i+j+k-9,1).merge(table.cell(i+j+k-9,2))
        table.cell(i+j+k-9,1).text=sheet_c.row_values(i-1)[11]

        table.cell(i+j+k-9,3).text=u'设计日期'
        table.cell(i+j+k-9,4).text=sheet_c.row_values(i-1)[12]


        document.add_page_break()
    else:
        add_row = table.add_row().cells
        table.cell(i+j+k,0).text=row_data[6]
        table.cell(i+j+k,1).text=row_data[7]
        table.cell(i+j+k,2).text=row_data[8]
        table.cell(i+j+k,3).text=row_data[9]
        table.cell(i+j+k,4).text=row_data[10]


    old=newv

document.add_page_break()
document.save('111.docx')
