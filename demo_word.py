from docx import Document

document = Document()

# 添加表
table = document.add_table(rows=2, cols=2,style='Light Grid')
cell = table.cell(0, 1)

cell.text = 'parrot, possibly dead'


table.add_row()

document.save('33.docx')